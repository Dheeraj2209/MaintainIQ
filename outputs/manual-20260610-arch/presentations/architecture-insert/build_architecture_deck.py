from pathlib import Path
import math
import shutil
import textwrap

from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Inches


ROOT = Path(r"C:\Users\LENOVO\OneDrive\Documents\IOT_idea_presentation")
SRC = ROOT / "Smart_Predictive_Maintenance_Review.pptx"
OUT = ROOT / "Smart_Predictive_Maintenance_Review_With_Architecture_Diagrams.pptx"
WORK = ROOT / "outputs" / "manual-20260610-arch" / "presentations" / "architecture-insert"
ASSETS = WORK / "diagram-assets"

W, H = 2400, 1350

COLORS = {
    "bg": (248, 251, 255),
    "band": (242, 246, 252),
    "ink": (22, 32, 50),
    "muted": (91, 103, 122),
    "line": (213, 223, 235),
    "blue": (37, 99, 235),
    "blue_light": (219, 234, 254),
    "teal": (15, 118, 110),
    "teal_light": (204, 251, 241),
    "green": (22, 128, 61),
    "green_light": (220, 252, 231),
    "red": (220, 38, 38),
    "red_light": (254, 226, 226),
    "orange": (194, 65, 12),
    "orange_light": (255, 237, 213),
    "purple": (124, 58, 237),
    "purple_light": (237, 233, 254),
    "white": (255, 255, 255),
}


def font(name="segoeui.ttf", size=32):
    return ImageFont.truetype(str(Path(r"C:\Windows\Fonts") / name), size=size)


F_TITLE = font("segoeuib.ttf", 48)
F_H2 = font("segoeuib.ttf", 30)
F_H3 = font("segoeuib.ttf", 24)
F_BODY = font("segoeui.ttf", 21)
F_SMALL = font("segoeui.ttf", 17)
F_TINY = font("segoeui.ttf", 15)
F_BOLD = font("segoeuib.ttf", 21)


def wrap_text(draw, text, max_width, fnt):
    words = text.split()
    lines = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if draw.textbbox((0, 0), candidate, font=fnt)[2] <= max_width or not current:
            current = candidate
        else:
            lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_wrapped(draw, xy, text, max_width, fnt=F_BODY, fill=None, line_gap=8):
    x, y = xy
    fill = fill or COLORS["ink"]
    for line in wrap_text(draw, text, max_width, fnt):
        draw.text((x, y), line, font=fnt, fill=fill)
        y += fnt.size + line_gap
    return y


def rounded(draw, box, fill, outline=None, width=2, radius=18):
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def arrow(draw, start, end, color, width=5, label=None, label_pos=0.5, label_fill=None):
    draw.line([start, end], fill=color, width=width)
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    size = 18
    p1 = (end[0] - size * math.cos(angle - math.pi / 7), end[1] - size * math.sin(angle - math.pi / 7))
    p2 = (end[0] - size * math.cos(angle + math.pi / 7), end[1] - size * math.sin(angle + math.pi / 7))
    draw.polygon([end, p1, p2], fill=color)
    if label:
        lx = start[0] + (end[0] - start[0]) * label_pos
        ly = start[1] + (end[1] - start[1]) * label_pos
        bbox = draw.textbbox((0, 0), label, font=F_TINY)
        pad = 8
        rounded(draw, (lx - bbox[2] / 2 - pad, ly - 20, lx + bbox[2] / 2 + pad, ly + 14), COLORS["bg"], None, 0, 7)
        draw.text((lx - bbox[2] / 2, ly - 18), label, font=F_TINY, fill=label_fill or COLORS["muted"])


def canvas(title, subtitle, chip):
    img = Image.new("RGB", (W, H), COLORS["bg"])
    draw = ImageDraw.Draw(img)
    draw.rectangle((0, 0, W, 116), fill=COLORS["white"])
    draw.text((64, 34), title, font=F_TITLE, fill=COLORS["ink"])
    draw.text((64, 88), subtitle, font=F_BODY, fill=COLORS["muted"])
    rounded(draw, (64, 150, 220, 194), COLORS["blue_light"], None, 0, 22)
    draw.text((94, 160), chip, font=F_SMALL, fill=COLORS["blue"])
    rounded(draw, (1780, 166, 2310, 1244), COLORS["white"], COLORS["line"], 2, 18)
    return img, draw


def notes(draw, items):
    y = 206
    for heading, body in items:
        draw.text((1818, y), heading, font=F_BOLD, fill=COLORS["ink"])
        y = draw_wrapped(draw, (1818, y + 34), body, 418, F_SMALL, COLORS["muted"], 6) + 58


def node(draw, box, title, body, accent="blue", icon=None):
    x1, y1, x2, y2 = box
    rounded(draw, box, COLORS["white"], COLORS["line"], 2, 14)
    draw.line((x1 + 2, y1 + 2, x2 - 2, y1 + 2), fill=COLORS[accent], width=7)
    if icon:
        ix, iy = x1 + 30, y1 + 44
        if icon == "db":
            draw.ellipse((ix, iy, ix + 52, iy + 20), outline=COLORS[accent], width=4)
            draw.rectangle((ix, iy + 10, ix + 52, iy + 58), outline=COLORS[accent], width=4)
            draw.ellipse((ix, iy + 48, ix + 52, iy + 68), outline=COLORS[accent], width=4)
        elif icon == "system":
            rounded(draw, (ix, iy, ix + 58, iy + 58), COLORS[f"{accent}_light"], COLORS[accent], 3, 8)
        elif icon == "alert":
            rounded(draw, (ix, iy + 12, ix + 58, iy + 48), COLORS[f"{accent}_light"], COLORS[accent], 3, 8)
            draw.text((ix + 19, iy + 8), "!", font=F_H3, fill=COLORS[accent])
        else:
            draw.ellipse((ix, iy, ix + 58, iy + 58), fill=COLORS[f"{accent}_light"], outline=COLORS[accent], width=3)
        tx = x1 + 110
    else:
        tx = x1 + 28
    draw.text((tx, y1 + 34), title, font=F_H3, fill=COLORS["ink"])
    draw_wrapped(draw, (tx, y1 + 70), body, max(50, x2 - tx - 28), F_SMALL, COLORS["muted"], 5)


def system_context(path):
    img, draw = canvas(
        "System Context Diagram (C4 Level 1)",
        "External users, monitored equipment, and the predictive maintenance system boundary",
        "C4 Level 1",
    )
    notes(draw, [
        ("Purpose", "Show how maintenance users, machine data, and optional datasets connect to the system."),
        ("System boundary", "The predictive maintenance system owns ingestion, analysis, alerts, storage, and dashboard outputs."),
        ("External parties", "Operators and supervisors consume actionable status, alerts, trends, and recommendations."),
        ("No UML actors", "External roles are shown as labeled context boxes, not use-case or actor notation."),
    ])
    rounded(draw, (520, 300, 1340, 1005), (255, 255, 255), COLORS["line"], 3, 24)
    draw.text((568, 340), "Smart Predictive Maintenance System", font=F_H2, fill=COLORS["ink"])
    node(draw, (590, 430, 920, 585), "Ingest data", "Sensor streams or stored vibration datasets", "blue", "system")
    node(draw, (965, 430, 1270, 585), "Analyze health", "Clean signal, extract features, classify condition", "purple", "system")
    node(draw, (590, 670, 920, 825), "Store evidence", "Raw readings, features, prediction history", "green", "db")
    node(draw, (965, 670, 1270, 825), "Notify users", "Severity, recommendation, dashboard status", "red", "alert")
    arrow(draw, (920, 508), (965, 508), COLORS["blue"], label="features")
    arrow(draw, (1118, 585), (1118, 670), COLORS["purple"], label="results")
    arrow(draw, (920, 748), (965, 748), COLORS["green"], label="trends")
    arrow(draw, (755, 585), (755, 670), COLORS["blue"], label="history")
    node(draw, (90, 390, 390, 540), "Monitored machine", "Motor, pump, fan, bearing or lab setup", "teal", "system")
    node(draw, (90, 640, 390, 790), "Vibration source", "Accelerometer hardware or simulated dataset", "orange", "db")
    node(draw, (1445, 380, 1715, 555), "Operator", "Views status and responds to alerts", "blue", "system")
    node(draw, (1445, 625, 1715, 800), "Supervisor", "Reviews trends and plans maintenance", "green", "system")
    arrow(draw, (390, 465), (590, 495), COLORS["teal"], label="vibration readings")
    arrow(draw, (390, 715), (590, 505), COLORS["orange"], label="demo data")
    arrow(draw, (1270, 500), (1445, 465), COLORS["blue"], label="dashboard")
    arrow(draw, (1270, 748), (1445, 715), COLORS["green"], label="trend evidence")
    arrow(draw, (1270, 545), (1445, 515), COLORS["red"], label="alerts")
    img.save(path)


def existing_vs_proposed(path):
    img, draw = canvas(
        "Existing Approach vs Proposed Approach",
        "How the proposed IoT + ML pipeline closes gaps in current maintenance practices",
        "Comparison",
    )
    notes(draw, [
        ("Purpose", "Position the proposed solution against reactive, scheduled, manual, threshold, and costly industrial options."),
        ("Reading order", "Left-to-right: current method, limitation, proposed capability."),
        ("Message", "The project adds early warning, condition-based action, and historical evidence without full industrial complexity."),
    ])
    headers = [("Current approaches", 90, COLORS["orange"]), ("Common limitations", 640, COLORS["red"]), ("Proposed capability", 1190, COLORS["green"])]
    for text, x, color in headers:
        rounded(draw, (x, 250, x + 440, 305), COLORS["white"], COLORS["line"], 2, 12)
        draw.line((x + 3, 250, x + 437, 250), fill=color, width=6)
        draw.text((x + 26, 264), text, font=F_H3, fill=COLORS["ink"])
    rows = [
        ("Reactive maintenance", "Fault detected late", "Early abnormal-pattern warning"),
        ("Scheduled maintenance", "Unnecessary service possible", "Condition-based maintenance"),
        ("Manual inspection", "Technician-dependent checks", "Trend evidence for decisions"),
        ("Threshold monitoring", "Misses gradual pattern shifts", "ML/rule health classification"),
        ("Industrial platforms", "High cost and complexity", "Accessible academic prototype"),
    ]
    y = 355
    for a, b, c in rows:
        node(draw, (90, y, 530, y + 120), a, "Current practice", "orange")
        node(draw, (640, y, 1080, y + 120), b, "Why it limits early maintenance", "red")
        node(draw, (1190, y, 1630, y + 120), c, "What the proposed system adds", "green")
        arrow(draw, (530, y + 60), (640, y + 60), COLORS["muted"], width=4)
        arrow(draw, (1080, y + 60), (1190, y + 60), COLORS["green"], width=4)
        y += 145
    rounded(draw, (90, 1115, 1630, 1215), COLORS["blue_light"], None, 0, 18)
    draw.text((128, 1142), "Conclusion: vibration data becomes a practical maintenance signal when it is cleaned, classified, stored, and converted into clear alerts.", font=F_H3, fill=COLORS["ink"])
    img.save(path)


def user_journey(path):
    img, draw = canvas(
        "User Journey Diagram",
        "Setup, monitoring, alert response, decision, and learning loop",
        "Journey",
    )
    notes(draw, [
        ("Purpose", "Explain the system from the perspective of students, operators, supervisors, and evaluators."),
        ("Flow", "The journey starts with setup, continues through monitoring and alert response, and ends in evidence-backed improvement."),
        ("Value", "Users receive understandable status and recommendations instead of raw ML complexity."),
    ])
    sections = [
        ("Setup", COLORS["teal"], 305, [("Attach sensor", "Student developer"), ("Configure ID", "Student developer"), ("Start collection", "Student developer")]),
        ("Monitoring", COLORS["blue"], 305, [("Collect readings", "System"), ("Analyze health", "System"), ("Display status", "Operator")]),
        ("Abnormal condition", COLORS["red"], 330, [("Detect pattern", "System"), ("Generate alert", "System"), ("Inspect machine", "Operator")]),
        ("Decision", COLORS["orange"], 260, [("Schedule maintenance", "Supervisor"), ("Confirm action", "Operator")]),
        ("Learning", COLORS["purple"], 260, [("Review trends", "Supervisor"), ("Add samples", "Student developer")]),
    ]
    x = 70
    top = 300
    columns = []
    for title, color, width, steps in sections:
        columns.append((x, width))
        rounded(draw, (x, top, x + width, 1065), COLORS["white"], COLORS["line"], 2, 18)
        draw.rectangle((x, top, x + width, top + 64), fill=color)
        draw.text((x + 24, top + 18), title, font=F_H3, fill=COLORS["white"])
        sy = top + 110
        previous_card = None
        for idx, (step, role) in enumerate(steps, 1):
            rounded(draw, (x + 28, sy, x + width - 28, sy + 110), (250, 252, 255), COLORS["line"], 2, 14)
            draw.ellipse((x + 48, sy + 34, x + 88, sy + 74), fill=color)
            draw.text((x + 61, sy + 39), str(idx), font=F_SMALL, fill=COLORS["white"])
            title_bottom = draw_wrapped(draw, (x + 104, sy + 20), step, width - 140, F_BOLD, COLORS["ink"], 0)
            role_y = max(sy + 58, title_bottom + 4)
            draw.text((x + 104, role_y), role, font=F_SMALL, fill=COLORS["muted"])
            if previous_card:
                arrow(draw, (x + width / 2, previous_card + 110), (x + width / 2, sy), color, width=4)
            previous_card = sy
            sy += 160
        x += width + 25
    for (x1, w1), (x2, _w2) in zip(columns, columns[1:]):
        arrow(draw, (x1 + w1 + 7, top + 32), (x2 - 7, top + 32), COLORS["blue"], width=4)
    img.save(path)


def end_to_end_sequence(path):
    img, draw = canvas(
        "End-to-End Sequence Diagram",
        "Typical runtime workflow from vibration signal to dashboard alert",
        "Sequence",
    )
    notes(draw, [
        ("Purpose", "Show the runtime chain that connects sensor readings, prediction, storage, alerting, and user action."),
        ("Key branch", "Abnormal output creates an alert; healthy output refreshes normal dashboard status."),
        ("Readable scope", "This is an end-to-end sequence, not a UML use-case or actor diagram."),
    ])
    participants = [
        ("Machine", 95, "teal"),
        ("Sensor", 285, "teal"),
        ("Gateway", 475, "blue"),
        ("Pipeline", 690, "purple"),
        ("Model", 905, "purple"),
        ("Storage", 1110, "green"),
        ("Alert", 1325, "red"),
        ("Dashboard", 1530, "blue"),
    ]
    y0, y1 = 250, 1165
    for name, x, accent in participants:
        rounded(draw, (x - 80, y0, x + 80, y0 + 58), COLORS["white"], COLORS["line"], 2, 12)
        draw.line((x - 77, y0, x + 77, y0), fill=COLORS[accent], width=6)
        tw = draw.textbbox((0, 0), name, font=F_BOLD)[2]
        draw.text((x - tw / 2, y0 + 18), name, font=F_BOLD, fill=COLORS["ink"])
        draw.line((x, y0 + 75, x, y1), fill=(198, 210, 225), width=3)
    events = [
        (340, 95, 285, "Produce vibration signal"),
        (420, 285, 475, "Send raw reading"),
        (500, 475, 690, "Transmit timestamped batch"),
        (580, 690, 690, "Clean, normalize, window"),
        (660, 690, 690, "Extract RMS, peak, mean, std"),
        (740, 690, 905, "Submit feature vector"),
        (820, 905, 690, "Return health state + confidence"),
        (900, 690, 1110, "Store features and result"),
    ]
    for y, x_from, x_to, label in events:
        if x_from == x_to:
            rounded(draw, (x_from + 16, y - 24, x_from + 265, y + 24), COLORS["white"], COLORS["line"], 2, 10)
            draw.text((x_from + 28, y - 13), label, font=F_TINY, fill=COLORS["muted"])
        else:
            arrow(draw, (x_from, y), (x_to, y), COLORS["blue"] if x_to > x_from else COLORS["muted"], width=4, label=label, label_pos=0.5)
    rounded(draw, (640, 950, 1665, 1095), (255, 249, 242), COLORS["orange"], 2, 14)
    draw.text((665, 972), "alt result is degrading or faulty", font=F_BOLD, fill=COLORS["orange"])
    arrow(draw, (690, 1030), (1325, 1030), COLORS["red"], width=4, label="request alert")
    arrow(draw, (1325, 1070), (1110, 1070), COLORS["red"], width=4, label="store alert")
    arrow(draw, (1325, 1110), (1530, 1110), COLORS["red"], width=4, label="push severity + recommendation")
    rounded(draw, (640, 1130, 1665, 1208), (239, 246, 255), COLORS["blue"], 2, 14)
    draw.text((665, 1152), "else healthy: storage provides latest status and dashboard refreshes normal trend", font=F_BOLD, fill=COLORS["blue"])
    img.save(path)


def normalize_existing(src, dst):
    im = Image.open(src).convert("RGB")
    im = im.resize((W, H), Image.Resampling.LANCZOS)
    im.save(dst, quality=95)


def build_pptx(diagrams):
    shutil.copyfile(SRC, OUT)
    prs = Presentation(str(OUT))
    original_ids = list(prs.slides._sldIdLst)
    blank = prs.slide_layouts[0]
    new_ids = {}
    for key, image_path in diagrams:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(str(image_path), 0, 0, width=prs.slide_width, height=prs.slide_height)
        new_ids[key] = prs.slides._sldIdLst[-1]

    final = []
    # Original slide indexes are zero-based here.
    final.extend(original_ids[0:3])
    final.append(new_ids["existing_vs_proposed"])
    final.extend(original_ids[3:4])
    final.append(new_ids["system_context"])
    final.extend(original_ids[4:5])
    final.append(new_ids["container"])
    final.append(new_ids["component"])
    final.extend(original_ids[5:7])
    final.append(new_ids["user_journey"])
    final.extend(original_ids[7:8])
    final.append(new_ids["sequence"])
    final.extend(original_ids[8:])

    sldIdLst = prs.slides._sldIdLst
    for elem in list(sldIdLst):
        sldIdLst.remove(elem)
    for elem in final:
        sldIdLst.append(elem)
    prs.save(str(OUT))


def main():
    ASSETS.mkdir(parents=True, exist_ok=True)
    generated = {
        "system_context": ASSETS / "system_context_c4_level_1.png",
        "existing_vs_proposed": ASSETS / "existing_vs_proposed.png",
        "user_journey": ASSETS / "user_journey.png",
        "sequence": ASSETS / "end_to_end_sequence.png",
        "container": ASSETS / "container_architecture_c4_level_2.png",
        "component": ASSETS / "component_architecture_c4_level_3.png",
    }
    system_context(generated["system_context"])
    existing_vs_proposed(generated["existing_vs_proposed"])
    user_journey(generated["user_journey"])
    end_to_end_sequence(generated["sequence"])
    normalize_existing(ROOT / "outputs" / "figma_container_preview.png", generated["container"])
    normalize_existing(ROOT / "outputs" / "figma_component_preview.png", generated["component"])
    build_pptx([
        ("existing_vs_proposed", generated["existing_vs_proposed"]),
        ("system_context", generated["system_context"]),
        ("container", generated["container"]),
        ("component", generated["component"]),
        ("user_journey", generated["user_journey"]),
        ("sequence", generated["sequence"]),
    ])
    print(OUT)


if __name__ == "__main__":
    main()
